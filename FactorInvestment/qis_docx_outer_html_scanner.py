import os
import re
from docx import Document
from pathlib import Path
from typing import List, Union


def extract_tickers_from_docx(file_path: Union[str, Path]) -> List[str]:

    """
    Extracts unique ticker symbols from 'data-rowkey' attributes in a Word document.

    :param file_path: Path to the .docx file on the computer.
    :return: List of (unique) extracted tickers, or an empty list if file is invalid.
    """

    # 1. Check if the file exists and is a valid file
    if not os.path.exists(file_path):
        print(f"Error: The path '{file_path}' does not exist.")
        return []

    if not os.path.isfile(file_path):
        print(f"Error: The path '{file_path}' is a directory, not a file.")
        return []

    # Regex pattern to match: data-rowkey="...:...
    # It captures everything after the colon (:) up to the closing quotation mark (")
    pattern = re.compile(r'data-rowkey="[^"]*?:([^"]+)"')
    tickers = []


    try:
        # Load the document
        doc = Document(file_path)

        # Scan all paragraphs in the document
        for paragraph in doc.paragraphs:
            if "data-rowkey" in paragraph.text:
                matches = pattern.findall(paragraph.text)
                tickers.extend(matches)

        # Scan all tables (in case the raw HTML content is inside table cells)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "data-rowkey" in cell.text:
                        matches = pattern.findall(cell.text)
                        tickers.extend(matches)

        # Remove duplicates while maintaining order
        tickers = list(dict.fromkeys(tickers))

        print(f"Extracted {len(tickers)} unique tickers.")
        return tickers

    except Exception as e:
        print(f"Error reading document: {e}")
        return []


if __name__ == "__main__":
    target_path = ""
    extracted_tickers = extract_tickers_from_docx(target_path)
    print(extracted_tickers)

